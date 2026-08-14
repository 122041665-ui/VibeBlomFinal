from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = "Infografia_Cuadro_Comparativo_POO.docx"

NAVY = "12324A"
BLUE = "146C94"
TEAL = "168C83"
PALE_BLUE = "EAF4F8"
PALE_GREEN = "EAF7F3"
LIGHT = "F7FAFC"
WHITE = "FFFFFF"
INK = "18323F"
MUTED = "526975"
BORDER = "B9D4DD"

rows = [
    ("◇  CLASE", "Es una plantilla que define las características y acciones que tendrán determinados objetos.", "Es el modelo general.", "La clase Automóvil establece que todos los automóviles tienen marca, color y velocidad."),
    ("●  OBJETO", "Es un elemento creado a partir de una clase. Tiene valores propios y puede ejecutar las acciones definidas en ella.", "Es un ejemplo específico creado a partir de la clase.", "Un automóvil Nissan de color rojo es un objeto de la clase Automóvil."),
    ("≡  ATRIBUTOS", "Son las características o datos que describen a una clase y a sus objetos.", "Representan lo que un objeto tiene o es.", "Un celular tiene marca, modelo, color y nivel de batería."),
    ("⚙  MÉTODOS", "Son las acciones o funciones que puede realizar un objeto.", "Representan lo que un objeto puede hacer.", "Un celular puede llamar, enviar mensajes, tomar fotografías o apagarse."),
    ("↳  HERENCIA", "Permite crear una clase nueva usando las características y métodos de otra clase existente.", "Sirve para reutilizar elementos de una clase.", "Perro y Gato pueden heredar de Animal características como nombre y edad."),
    ("⇄  POLIMORFISMO", "Permite que un mismo método se comporte de distintas maneras según el objeto que lo utilice.", "Permite realizar una misma acción de distintas formas.", "Perro y Gato usan hacerSonido(): el perro ladra y el gato maúlla."),
    ("＋  SOBRECARGA", "Consiste en crear varios métodos con el mismo nombre, pero con distinta cantidad o tipo de parámetros.", "Normalmente ocurre dentro de una misma clase.", "pagar() puede utilizarse para pagar con efectivo, tarjeta o transferencia."),
    ("↻  SOBRESCRITURA", "Ocurre cuando una clase hija modifica un método heredado para darle un comportamiento diferente.", "Reemplaza el funcionamiento de un método heredado.", "Animal tiene moverse(), pero un pez lo modifica para nadar y un ave para volar."),
    ("▣  ENCAPSULAMIENTO", "Protege los datos internos de un objeto y permite acceder a ellos solo mediante métodos autorizados.", "Evita modificar la información directa o incorrectamente.", "El saldo bancario no se modifica directamente: se usan depositar y retirar."),
    ("⛓  RELACIONES ENTRE OBJETOS", "Son conexiones entre dos o más objetos para colaborar en un sistema: asociación, agregación o composición.", "Permiten compartir información y trabajar juntos.", "Un estudiante pertenece a un grupo; un cliente realiza pedidos; un automóvil tiene un motor."),
    ("✦  INSTANCIAMIENTO DE OBJETOS", "Es el proceso de crear un objeto concreto a partir de una clase.", "Convierte una plantilla general en un elemento utilizable en el programa.", "Desde la clase Persona se crea el objeto Rafael, con edad y nombre propios."),
]

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)

def set_cell_margins(cell, top=55, start=90, bottom=55, end=90):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v)); node.set(qn("w:type"), "dxa")

def borders(table):
    tblPr = table._tbl.tblPr
    b = tblPr.first_child_found_in("w:tblBorders")
    if b is None:
        b = OxmlElement("w:tblBorders"); tblPr.append(b)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "5")
        el.set(qn("w:color"), BORDER)
        b.append(el)

def font(run, size, color=INK, bold=False, name="Aptos"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size); run.font.color.rgb = RGBColor.from_string(color); run.bold = bold

def add_rich_text(cell, text, size=6.65, color=INK):
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 0.92
    keywords = ["Automóvil", "tiene o es", "puede hacer", "Perro", "Gato", "Animal", "hacerSonido()", "pagar()", "moverse()", "Persona", "Rafael", "asociación, agregación o composición"]
    pos = 0
    matches = []
    for kw in keywords:
        start = text.find(kw)
        if start >= 0: matches.append((start, start + len(kw)))
    matches.sort()
    for start, end in matches:
        if start < pos: continue
        if start > pos: font(p.add_run(text[pos:start]), size, color)
        font(p.add_run(text[start:end]), size, BLUE, True)
        pos = end
    if pos < len(text): font(p.add_run(text[pos:]), size, color)

doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width = Inches(13.333)
sec.page_height = Inches(7.5)
sec.top_margin = Inches(0.24); sec.bottom_margin = Inches(0.20)
sec.left_margin = Inches(0.28); sec.right_margin = Inches(0.28)
sec.header_distance = Inches(0.1); sec.footer_distance = Inches(0.1)

normal = doc.styles["Normal"]
normal.font.name = "Aptos"; normal.font.size = Pt(7)
normal.paragraph_format.space_after = Pt(0)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(1)
p.paragraph_format.line_spacing = 0.9
r = p.add_run("PROGRAMACIÓN ORIENTADA A OBJETOS")
font(r, 17.5, NAVY, True)
r2 = p.add_run("   ·   CUADRO COMPARATIVO")
font(r2, 10.5, TEAL, True)

p2 = doc.add_paragraph()
p2.paragraph_format.space_after = Pt(5); p2.paragraph_format.line_spacing = 0.9
font(p2.add_run("Conceptos esenciales explicados con diferencias clave y ejemplos cotidianos"), 7.5, MUTED)

table = doc.add_table(rows=1, cols=4)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
widths = [1.77, 4.10, 3.04, 3.78]
headers = ["CONCEPTO", "DEFINICIÓN Y CARACTERÍSTICAS", "DIFERENCIA PRINCIPAL", "EJEMPLO COTIDIANO"]
for i, (c, w, label) in enumerate(zip(table.rows[0].cells, widths, headers)):
    c.width = Inches(w); shade(c, NAVY); c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(c, 65, 95, 65, 95)
    p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(0)
    font(p.add_run(label), 7.25, WHITE, True)

for idx, rowdata in enumerate(rows):
    cells = table.add_row().cells
    fill = WHITE if idx % 2 == 0 else LIGHT
    for j, (c, w) in enumerate(zip(cells, widths)):
        c.width = Inches(w); c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(c, 48, 90, 48, 90)
        shade(c, PALE_GREEN if j == 0 and idx % 2 == 0 else (PALE_BLUE if j == 0 else fill))
        if j == 0:
            p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 0.9
            font(p.add_run(rowdata[j]), 6.9 if len(rowdata[j]) < 24 else 6.25, TEAL, True)
        else:
            add_rich_text(c, rowdata[j])

borders(table)

# Fixed table geometry in DXA, including indent matching left cell margin.
tblPr = table._tbl.tblPr
tblW = tblPr.first_child_found_in("w:tblW")
tblW.set(qn("w:w"), str(sum(round(w * 1440) for w in widths))); tblW.set(qn("w:type"), "dxa")
tblLayout = OxmlElement("w:tblLayout"); tblLayout.set(qn("w:type"), "fixed"); tblPr.append(tblLayout)
tblInd = OxmlElement("w:tblInd"); tblInd.set(qn("w:w"), "90"); tblInd.set(qn("w:type"), "dxa"); tblPr.append(tblInd)
grid = table._tbl.tblGrid
for gc, w in zip(grid.gridCol_lst, widths): gc.set(qn("w:w"), str(round(w * 1440)))
for tr in table.rows:
    for c, w in zip(tr.cells, widths):
        tcW = c._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
        tcW.set(qn("w:w"), str(round(w * 1440))); tcW.set(qn("w:type"), "dxa")

doc.save(OUT)
print(OUT)
